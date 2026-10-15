"""
================================================================================
DATA DICTIONARY GENERATOR SCRIPT
Generates both Data_Dictionary.docx and Data_Dictionary.md for Lesson 7 / Lab 4.
Group 10: Real Estate Property Listing & Regional Market Value Warehouse
================================================================================
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
DOCX_PATH = os.path.join(CURRENT_DIR, "Data_Dictionary.docx")
MD_PATH = os.path.join(CURRENT_DIR, "Data_Dictionary.md")


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_custom_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(26, 54, 93)  # Navy
        run.bold = True
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(43, 84, 126)  # Slate Blue
        run.bold = True
    elif level == 3:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(60, 60, 60)
        run.bold = True
    return h


def create_styled_table(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1A365D")  # Dark Blue
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
            run.font.name = "Arial"

    # Data Rows
    for r_idx, row_values in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F7FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(8.5)
                run.font.name = "Arial"

    # Set Widths if provided
    if col_widths:
        for row in table.rows:
            for c_idx, width in enumerate(col_widths):
                row.cells[c_idx].width = width

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


# Data definitions for Core Tables
TABLES_METADATA = {
    "dim_date": {
        "description": "Temporal dimension providing calendar attributes, quarterly partitioning, and weekend flags.",
        "headers": ["Column Name", "Data Type", "Constraints", "Key Type", "Source Mapping", "Business Definition"],
        "widths": [Inches(1.2), Inches(0.8), Inches(1.1), Inches(0.8), Inches(1.1), Inches(2.2)],
        "rows": [
            ["date_id", "INTEGER", "PRIMARY KEY", "PK", "date (derived)", "Surrogate integer date key formatted as YYYYMMDD"],
            ["full_date", "TEXT", "NOT NULL", "Attribute", "date", "Normalized ISO calendar date string (YYYY-MM-DD)"],
            ["year", "INTEGER", "NOT NULL, 1900-2100", "Attribute", "date (extracted)", "Calendar year of property sale transaction"],
            ["quarter", "INTEGER", "NOT NULL, 1-4", "Attribute", "date (extracted)", "Fiscal/calendar quarter (Q1 through Q4)"],
            ["month", "INTEGER", "NOT NULL, 1-12", "Attribute", "date (extracted)", "Numerical calendar month (1 through 12)"],
            ["month_name", "TEXT", "NOT NULL", "Attribute", "date (extracted)", "Full textual month designation (e.g., 'May')"],
            ["day", "INTEGER", "NOT NULL, 1-31", "Attribute", "date (extracted)", "Calendar day of month (1 through 31)"],
            ["day_of_week", "INTEGER", "NOT NULL, 0-6", "Attribute", "date (extracted)", "Day of week index (0=Monday, 6=Sunday)"],
            ["day_name", "TEXT", "NOT NULL", "Attribute", "date (extracted)", "Full textual day of week name (e.g., 'Friday')"],
            ["is_weekend", "INTEGER", "NOT NULL, 0 or 1", "Attribute", "date (derived)", "Binary boolean flag: 1 if Saturday/Sunday, else 0"]
        ]
    },
    "dim_location": {
        "description": "Geographical dimension capturing postal zones and precise GPS centroids.",
        "headers": ["Column Name", "Data Type", "Constraints", "Key Type", "Source Mapping", "Business Definition"],
        "widths": [Inches(1.2), Inches(0.8), Inches(1.1), Inches(0.8), Inches(1.1), Inches(2.2)],
        "rows": [
            ["location_id", "INTEGER", "PRIMARY KEY AUTO", "PK", "Derived ID", "Surrogate geographic location identifier"],
            ["zipcode", "TEXT", "NOT NULL, 5-digits", "Attribute", "zipcode", "Standard 5-digit United States postal code"],
            ["latitude", "REAL", "NOT NULL, [-90, 90]", "Attribute", "lat", "Geodetic latitude coordinate in decimal degrees"],
            ["longitude", "REAL", "NOT NULL, [-180, 180]", "Attribute", "long", "Geodetic longitude coordinate in decimal degrees"]
        ]
    },
    "dim_property": {
        "description": "Physical asset dimension holding structural vintage and neighborhood reference footprints.",
        "headers": ["Column Name", "Data Type", "Constraints", "Key Type", "Source Mapping", "Business Definition"],
        "widths": [Inches(1.2), Inches(0.8), Inches(1.1), Inches(0.8), Inches(1.1), Inches(2.2)],
        "rows": [
            ["property_id", "INTEGER", "PRIMARY KEY", "PK", "id", "Unique public parcel identification number"],
            ["yr_built", "INTEGER", "NOT NULL, >= 1800", "Attribute", "yr_built", "Original construction year of the main structure"],
            ["yr_renovated", "INTEGER", "NOT NULL, >= 0", "Attribute", "yr_renovated", "Most recent major renovation year (0 if never renovated)"],
            ["is_renovated", "INTEGER", "NOT NULL, 0 or 1", "Attribute", "Derived flag", "Binary flag indicating structural restoration (1=Yes, 0=No)"],
            ["sqft_living15", "INTEGER", "NOT NULL, > 0", "Attribute", "sqft_living15", "Average interior living space of the nearest 15 neighboring homes"],
            ["sqft_lot15", "INTEGER", "NOT NULL, > 0", "Attribute", "sqft_lot15", "Average parcel land area of the nearest 15 neighboring homes"]
        ]
    },
    "dim_condition_grade": {
        "description": "Structural assessment dimension storing building construction quality and site vistas.",
        "headers": ["Column Name", "Data Type", "Constraints", "Key Type", "Source Mapping", "Business Definition"],
        "widths": [Inches(1.2), Inches(0.8), Inches(1.1), Inches(0.8), Inches(1.1), Inches(2.2)],
        "rows": [
            ["grade_condition_id", "INTEGER", "PRIMARY KEY AUTO", "PK", "Derived ID", "Surrogate structural rating identifier"],
            ["grade", "INTEGER", "NOT NULL, 1-13", "Attribute", "grade", "Official King County building construction grade (1=Low, 13=Mansion)"],
            ["condition", "INTEGER", "NOT NULL, 1-5", "Attribute", "condition", "Physical maintenance condition rating (1=Poor, 5=Excellent)"],
            ["view", "INTEGER", "NOT NULL, 0-4", "Attribute", "view", "Visual scenic view quality score (0=None, 4=Panoramic)"],
            ["waterfront", "INTEGER", "NOT NULL, 0 or 1", "Attribute", "waterfront", "Binary indicator: 1 if parcel fronts a body of water, else 0"]
        ]
    },
    "fact_property_sales": {
        "description": "Central transactional fact table recording completed real estate sales and metrics.",
        "headers": ["Column Name", "Data Type", "Constraints", "Key Type", "Source Mapping", "Business Definition"],
        "widths": [Inches(1.2), Inches(0.8), Inches(1.1), Inches(0.8), Inches(1.1), Inches(2.2)],
        "rows": [
            ["sale_id", "INTEGER", "PRIMARY KEY AUTO", "PK", "Auto Seq", "Unique transactional surrogate sales identifier"],
            ["property_id", "INTEGER", "NOT NULL", "FK", "id", "Foreign key linking to dim_property(property_id)"],
            ["date_id", "INTEGER", "NOT NULL", "FK", "Derived date", "Foreign key linking to dim_date(date_id)"],
            ["location_id", "INTEGER", "NOT NULL", "FK", "Derived loc", "Foreign key linking to dim_location(location_id)"],
            ["grade_condition_id", "INTEGER", "NOT NULL", "FK", "Derived grade", "Foreign key linking to dim_condition_grade(grade_condition_id)"],
            ["price", "REAL", "NOT NULL, CHECK > 0", "Measure", "price", "Total closing sales transaction price in USD ($)"],
            ["sqft_living", "INTEGER", "NOT NULL, CHECK > 0", "Measure", "sqft_living", "Interior living area of the property in square feet"],
            ["sqft_lot", "INTEGER", "NOT NULL, CHECK > 0", "Measure", "sqft_lot", "Total land parcel lot area in square feet"],
            ["price_per_sqft", "REAL", "NOT NULL, CHECK > 0", "Measure", "Calculated", "Unit valuation metric: price divided by sqft_living ($/sqft)"],
            ["bedrooms", "INTEGER", "NOT NULL, CHECK >= 0", "Attribute", "bedrooms", "Total count of bedrooms in the dwelling"],
            ["bathrooms", "REAL", "NOT NULL, CHECK >= 0", "Attribute", "bathrooms", "Total count of bathrooms (quarter/half increments allowed)"],
            ["floors", "REAL", "NOT NULL, CHECK > 0", "Attribute", "floors", "Number of vertical finished floors / stories"],
            ["sqft_above", "INTEGER", "NOT NULL, CHECK >= 0", "Measure", "sqft_above", "Interior living space located strictly above grade level"],
            ["sqft_basement", "INTEGER", "NOT NULL, CHECK >= 0", "Measure", "sqft_basement", "Subterranean basement interior area in square feet"],
            ["has_basement", "INTEGER", "NOT NULL, 0 or 1", "Attribute", "Derived flag", "Binary flag indicating basement existence (1=Yes, 0=No)"]
        ]
    }
}

DATA_MART_TABLES = {
    "monthly_property_valuation_mart": {
        "description": "Downstream aggregate table summarizing chronological market pricing and volume.",
        "headers": ["Column Name", "Data Type", "Constraints", "Business Definition"],
        "widths": [Inches(1.8), Inches(1.0), Inches(1.4), Inches(3.0)],
        "rows": [
            ["year", "INTEGER", "NOT NULL", "Calendar transaction year"],
            ["month", "INTEGER", "NOT NULL", "Calendar transaction month (1-12)"],
            ["month_name", "TEXT", "NOT NULL", "Full month name (e.g., 'May')"],
            ["total_transactions", "INTEGER", "NOT NULL", "Total number of closed property sales"],
            ["total_sales_volume", "REAL", "NOT NULL", "Total dollar sum of closed transactions ($)"],
            ["avg_property_price", "REAL", "NOT NULL", "Arithmetic mean property sale price ($)"],
            ["min_property_price", "REAL", "NOT NULL", "Minimum closed property price in the month ($)"],
            ["max_property_price", "REAL", "NOT NULL", "Maximum closed property price in the month ($)"],
            ["avg_price_per_sqft", "REAL", "NOT NULL", "Mean market unit price per square foot ($/sqft)"],
            ["total_sqft_living", "INTEGER", "NOT NULL", "Cumulative living square footage sold"],
            ["avg_sqft_living", "REAL", "NOT NULL", "Average interior living area of sold dwellings"]
        ]
    },
    "regional_zipcode_valuation_mart": {
        "description": "Downstream aggregate table partitioning real estate performance by postal zones.",
        "headers": ["Column Name", "Data Type", "Constraints", "Business Definition"],
        "widths": [Inches(1.8), Inches(1.0), Inches(1.4), Inches(3.0)],
        "rows": [
            ["zipcode", "TEXT", "PRIMARY / INDEXED", "5-digit regional postal code"],
            ["centroid_latitude", "REAL", "NOT NULL", "Average latitude centroid for the postal region"],
            ["centroid_longitude", "REAL", "NOT NULL", "Average longitude centroid for the postal region"],
            ["total_listings_sold", "INTEGER", "NOT NULL", "Volume of properties sold in the postal zone"],
            ["total_regional_sales_volume", "REAL", "NOT NULL", "Aggregate dollar sales volume in the zipcode ($)"],
            ["avg_market_value", "REAL", "NOT NULL", "Mean residential property market valuation ($)"],
            ["avg_price_per_sqft", "REAL", "NOT NULL", "Mean unit price per living square foot ($/sqft)"],
            ["avg_living_area_sqft", "REAL", "NOT NULL", "Mean interior living space in square feet"],
            ["waterfront_property_count", "INTEGER", "NOT NULL", "Number of waterfront properties sold in region"],
            ["avg_bedrooms", "REAL", "NOT NULL", "Mean bedroom count across regional homes"],
            ["avg_bathrooms", "REAL", "NOT NULL", "Mean bathroom count across regional homes"]
        ]
    },
    "structural_grade_valuation_mart": {
        "description": "Downstream aggregate table evaluating valuation premiums by architectural grade tiers.",
        "headers": ["Column Name", "Data Type", "Constraints", "Business Definition"],
        "widths": [Inches(1.8), Inches(1.0), Inches(1.4), Inches(3.0)],
        "rows": [
            ["grade", "INTEGER", "NOT NULL", "King County construction grade rating (1-13)"],
            ["condition", "INTEGER", "NOT NULL", "Maintenance condition score (1-5)"],
            ["market_tier", "TEXT", "NOT NULL", "Categorical tier (Economy, Mid-Market, Executive, Luxury Estate)"],
            ["property_count", "INTEGER", "NOT NULL", "Number of properties belonging to this tier"],
            ["aggregate_tier_valuation", "REAL", "NOT NULL", "Total financial sales volume within tier ($)"],
            ["avg_valuation", "REAL", "NOT NULL", "Mean transaction sales price in tier ($)"],
            ["avg_price_per_sqft", "REAL", "NOT NULL", "Mean unit price per square foot in tier ($/sqft)"],
            ["avg_sqft_living", "REAL", "NOT NULL", "Mean interior living area in square feet"],
            ["avg_lot_size", "REAL", "NOT NULL", "Mean land parcel size in square feet"]
        ]
    }
}

INDEX_METADATA = [
    ["idx_fact_date", "fact_property_sales", "date_id", "Single B-Tree", "Optimizes temporal joins to dim_date and monthly filtering"],
    ["idx_fact_location", "fact_property_sales", "location_id", "Single B-Tree", "Accelerates regional spatial joins and postal aggregations"],
    ["idx_fact_grade", "fact_property_sales", "grade_condition_id", "Single B-Tree", "Accelerates joins on property construction ratings"],
    ["idx_fact_property", "fact_property_sales", "property_id", "Single B-Tree", "Speeds up parcel lookup and property history tracing"],
    ["idx_dim_loc_zip", "dim_location", "zipcode", "Single B-Tree", "Accelerates zipcode filtering and spatial grouping"],
    ["idx_dim_date_ym", "dim_date", "year, month", "Compound B-Tree", "Enables index skip scans on year/month time-series queries"],
    ["idx_fact_date_price", "fact_property_sales", "date_id, price", "Compound B-Tree", "Eliminates table scans when querying date ranges with price thresholds"],
    ["idx_fact_loc_price_sqft", "fact_property_sales", "location_id, price, sqft_living", "Compound Covering", "Directly satisfies regional valuation queries without table heap reads"],
    ["idx_fact_grade_price", "fact_property_sales", "grade_condition_id, price", "Compound B-Tree", "Replaces full table scans with index searches for luxury grade filtering"],
    ["idx_fact_bed_price", "fact_property_sales", "bedrooms, price", "Compound B-Tree", "Accelerates bedroom-level market segment aggregations"],
    ["2dsphere (Mongo)", "property_features", "neighborhood_context.location", "2dsphere Spatial", "Enables geospatial indexing and radius queries ($near, $geoWithin)"]
]


def generate_docx():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Title Banner
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("TECHNICAL DATA DICTIONARY")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(26, 54, 93)

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle_p.add_run("Data Warehousing (LEC/LAB) — Lesson 7: Performance Tuning & Data Mart Creation\n"
                                 "Group 10: Real Estate Property Listing & Regional Market Value Warehouse\n"
                                 "Instructor: Prof. Rob Malitao | Date: October 2026")
    sub_run.font.size = Pt(10)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 1: Overview
    add_custom_heading(doc, "1. Executive Summary & Warehouse Architecture", level=1)
    p_desc = doc.add_paragraph()
    p_desc.add_run(
        "This technical data dictionary comprehensively documents the data architecture, relational star schema, "
        "data governance constraints, B-Tree index specifications, MongoDB document store schema, and the "
        "materialized Real Estate Departmental Data Mart for Group 10. The warehouse ingests King County residential "
        "property transactions (21,613 records) and enforces dual-engine analytical processing across SQLite and MongoDB."
    )

    # Section 2: Core Relational Star Schema
    add_custom_heading(doc, "2. Core Relational Star Schema Specifications (SQLite)", level=1)
    for tbl_name, meta in TABLES_METADATA.items():
        add_custom_heading(doc, f"Table: {tbl_name}", level=2)
        doc.add_paragraph(meta["description"])
        create_styled_table(doc, meta["headers"], meta["rows"], meta["widths"])

    # Section 3: Data Governance Constraints
    add_custom_heading(doc, "3. Data Governance & Relational Constraint Rules", level=1)
    p_gov = doc.add_paragraph()
    p_gov.add_run(
        "Data governance rules are enforced at the DDL table creation level to prevent dirty, malformed, or corrupt "
        "records from entering the analytical repository. The table below details active constraints:"
    )
    gov_headers = ["Rule Category", "Target Table", "Target Column(s)", "Enforced Constraint", "Business Rationale"]
    gov_widths = [Inches(1.2), Inches(1.3), Inches(1.2), Inches(1.6), Inches(1.9)]
    gov_rows = [
        ["Financial Integrity", "fact_property_sales", "price", "CHECK (price > 0.0)", "Guarantees closing sale prices are strictly positive."],
        ["Physical Integrity", "fact_property_sales", "sqft_living, sqft_lot", "CHECK (sqft > 0)", "Prevents zero or negative square footage anomalies."],
        ["Unit Rate Integrity", "fact_property_sales", "price_per_sqft", "CHECK (price_per_sqft > 0.0)", "Enforces calculated unit rate positivity."],
        ["Structural Bounds", "dim_condition_grade", "grade", "CHECK (grade BETWEEN 1 AND 13)", "Enforces King County building code boundaries."],
        ["Maintenance Bounds", "dim_condition_grade", "condition", "CHECK (condition BETWEEN 1 AND 5)", "Validates building condition inspection score."],
        ["Temporal Integrity", "dim_property", "yr_built", "CHECK (yr_built BETWEEN 1800 AND 2050)", "Blocks corrupted vintage years."],
        ["Spatial Bounds", "dim_location", "latitude, longitude", "CHECK (lat [-90,90], long [-180,180])", "Blocks out-of-range geodetic coordinates."],
        ["Referential Integrity", "fact_property_sales", "Foreign Keys", "NOT NULL, ON DELETE RESTRICT", "Guarantees no orphaned transactional fact rows exist."]
    ]
    create_styled_table(doc, gov_headers, gov_rows, gov_widths)

    # Section 4: MongoDB Document Store Schema
    add_custom_heading(doc, "4. Document Store Schema & $jsonSchema Governance (MongoDB)", level=1)
    p_mongo = doc.add_paragraph()
    p_mongo.add_run(
        "For non-relational nested analysis, documents in the 'property_architectural_features_governed' collection "
        "are strictly validated at ingestion time using MongoDB's $jsonSchema validator with validationLevel='strict' "
        "and validationAction='error'."
    )
    mongo_headers = ["BSON Field Path", "BSON Type", "Validation Rule", "Business Definition"]
    mongo_widths = [Inches(2.2), Inches(1.0), Inches(1.5), Inches(2.5)]
    mongo_rows = [
        ["_id / property_id", "int / long", "Required, Unique", "Unique numeric property identification."],
        ["architectural_specs.bedrooms", "int", "minimum: 0", "Non-negative bedroom count."],
        ["architectural_specs.bathrooms", "double / int", "minimum: 0.0", "Non-negative bathroom count."],
        ["architectural_specs.sqft_living", "int", "minimum: 1", "Strictly positive interior living area."],
        ["evaluation_metrics.grade", "int", "minimum: 1, maximum: 13", "Enforces valid architectural quality range."],
        ["evaluation_metrics.condition", "int", "minimum: 1, maximum: 5", "Enforces physical inspection rating scale."],
        ["neighborhood_context.zipcode", "string", "pattern: '^[0-9]{5}$'", "Enforces 5-digit US postal code regex."],
        ["neighborhood_context.location", "object", "GeoJSON Point [lng, lat]", "Spatial 2dsphere indexing and geo queries."],
        ["sales_history.price", "double / int", "minimum: 1.0", "Blocks negative or zero sales prices."]
    ]
    create_styled_table(doc, mongo_headers, mongo_rows, mongo_widths)

    # Section 5: Departmental Data Mart Schema
    add_custom_heading(doc, "5. Departmental Data Mart Schema (property_valuation_mart.db)", level=1)
    p_mart = doc.add_paragraph()
    p_mart.add_run(
        "To isolate business intelligence reporting from transactional load, the departmental data mart extracts "
        "and materializes pre-aggregated property valuation summaries into an isolated SQLite database:"
    )
    for tbl_name, meta in DATA_MART_TABLES.items():
        add_custom_heading(doc, f"Data Mart Table: {tbl_name}", level=2)
        doc.add_paragraph(meta["description"])
        create_styled_table(doc, meta["headers"], meta["rows"], meta["widths"])

    # Section 6: Performance Tuning & Index Catalog
    add_custom_heading(doc, "6. Index Catalog & Performance Tuning Audit", level=1)
    p_idx = doc.add_paragraph()
    p_idx.add_run(
        "To eliminate expensive full table scans (SCAN fact_property_sales across 21,613 records), custom single-column "
        "and compound B-Tree indexes were created on high-cardinality foreign keys and filter fields. "
        "Benchmark testing proved up to a 124x query speedup with SEARCH ... USING INDEX plans."
    )
    idx_headers = ["Index Name", "Target Table", "Indexed Columns", "Index Type", "Tuning & Query Plan Impact"]
    idx_widths = [Inches(1.5), Inches(1.3), Inches(1.5), Inches(1.1), Inches(1.8)]
    create_styled_table(doc, idx_headers, INDEX_METADATA, idx_widths)

    doc.save(DOCX_PATH)
    print(f"[DOCX GENERATION] Created {DOCX_PATH}")


def generate_markdown():
    lines = []
    lines.append("# Technical Data Dictionary Document")
    lines.append("**Data Warehousing (LEC/LAB) — Lesson 7: Performance Tuning & Data Mart Creation**  ")
    lines.append("**Group 10: Real Estate Property Listing & Regional Market Value Warehouse**  ")
    lines.append("**Instructor: Prof. Rob Malitao** | **Date: October 2026**\n")

    lines.append("## 1. Executive Summary & Warehouse Architecture")
    lines.append(
        "This data dictionary provides comprehensive metadata documentation for the Group 10 Real Estate Data Warehouse. "
        "It details the relational star schema (SQLite), strict data governance constraints, MongoDB document store schema, "
        "B-Tree performance indexes, and the materialized Real Estate Departmental Data Mart (`property_valuation_mart.db`).\n"
    )

    lines.append("## 2. Core Relational Star Schema (SQLite: `warehouse.db`)")
    for tbl_name, meta in TABLES_METADATA.items():
        lines.append(f"### Table: `{tbl_name}`")
        lines.append(f"*{meta['description']}*\n")
        lines.append("| " + " | ".join(meta["headers"]) + " |")
        lines.append("| " + " | ".join(["---"] * len(meta["headers"])) + " |")
        for r in meta["rows"]:
            lines.append("| " + " | ".join(str(x) for x in r) + " |")
        lines.append("")

    lines.append("## 3. Data Governance & Relational Constraint Rules")
    lines.append(
        "Strict DDL constraints enforced in `governed_warehouse.db` to prevent invalid or corrupted records:\n"
    )
    lines.append("| Rule Category | Target Table | Target Column(s) | Enforced Constraint | Business Rationale |")
    lines.append("| --- | --- | --- | --- | --- |")
    lines.append("| Financial Integrity | fact_property_sales | price | CHECK (price > 0.0) | Guarantees transaction price is strictly positive. |")
    lines.append("| Physical Integrity | fact_property_sales | sqft_living, sqft_lot | CHECK (sqft > 0) | Prevents zero or negative dwelling space. |")
    lines.append("| Unit Rate Integrity | fact_property_sales | price_per_sqft | CHECK (price_per_sqft > 0.0) | Enforces calculated valuation rate positivity. |")
    lines.append("| Structural Bounds | dim_condition_grade | grade | CHECK (grade BETWEEN 1 AND 13) | Enforces official King County grading codes. |")
    lines.append("| Maintenance Bounds | dim_condition_grade | condition | CHECK (condition BETWEEN 1 AND 5) | Validates inspection maintenance rating scale. |")
    lines.append("| Temporal Integrity | dim_property | yr_built | CHECK (yr_built BETWEEN 1800 AND 2050) | Blocks corrupted vintage years. |")
    lines.append("| Spatial Bounds | dim_location | latitude, longitude | CHECK (lat [-90,90], long [-180,180]) | Blocks out-of-range geodetic coordinates. |")
    lines.append("| Referential Integrity | fact_property_sales | Foreign Keys | NOT NULL, ON DELETE RESTRICT | Guarantees no orphaned fact rows. |\n")

    lines.append("## 4. Document Store Schema & $jsonSchema Governance (MongoDB)")
    lines.append(
        "Enforces strict document validation on `property_architectural_features_governed` with `validationLevel='strict'`:\n"
    )
    lines.append("| BSON Field Path | BSON Type | Validation Rule | Business Definition |")
    lines.append("| --- | --- | --- | --- |")
    lines.append("| `_id / property_id` | int / long | Required, Unique | Unique parcel identification. |")
    lines.append("| `architectural_specs.bedrooms` | int | minimum: 0 | Non-negative bedroom count. |")
    lines.append("| `architectural_specs.bathrooms` | double / int | minimum: 0.0 | Non-negative bathroom count. |")
    lines.append("| `architectural_specs.sqft_living` | int | minimum: 1 | Strictly positive interior living area. |")
    lines.append("| `evaluation_metrics.grade` | int | minimum: 1, maximum: 13 | Validates construction quality rating. |")
    lines.append("| `evaluation_metrics.condition` | int | minimum: 1, maximum: 5 | Validates maintenance condition. |")
    lines.append("| `neighborhood_context.zipcode` | string | pattern: `^[0-9]{5}$` | Enforces 5-digit US postal pattern. |")
    lines.append("| `neighborhood_context.location` | object | GeoJSON Point `[lng, lat]` | 2dsphere spatial indexing. |")
    lines.append("| `sales_history.price` | double / int | minimum: 1.0 | Blocks negative/zero closing prices. |\n")

    lines.append("## 5. Departmental Data Mart Schema (`property_valuation_mart.db`)")
    for tbl_name, meta in DATA_MART_TABLES.items():
        lines.append(f"### Data Mart Table: `{tbl_name}`")
        lines.append(f"*{meta['description']}*\n")
        lines.append("| " + " | ".join(meta["headers"]) + " |")
        lines.append("| " + " | ".join(["---"] * len(meta["headers"])) + " |")
        for r in meta["rows"]:
            lines.append("| " + " | ".join(str(x) for x in r) + " |")
        lines.append("")

    lines.append("## 6. Index Catalog & Performance Tuning Audit")
    lines.append("| Index Name | Target Table | Indexed Columns | Index Type | Tuning & Query Plan Impact |")
    lines.append("| --- | --- | --- | --- | --- |")
    for row in INDEX_METADATA:
        lines.append("| " + " | ".join(str(x) for x in row) + " |")
    lines.append("")

    with open(MD_PATH, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"[MARKDOWN GENERATION] Created {MD_PATH}")


if __name__ == "__main__":
    generate_docx()
    generate_markdown()
